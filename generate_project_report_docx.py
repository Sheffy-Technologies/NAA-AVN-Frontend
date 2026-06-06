from datetime import date
from pathlib import Path
from docx import Document
from docx.shared import Inches

root = Path('/Users/ram/projects/avn-frontend')
shots = root / 'screenshots'
out = root / 'AVN_Project_Report_With_Screenshots.docx'

doc = Document()
doc.add_heading('NAAVN DDBMS Frontend Project Report', 0)

doc.add_paragraph(f'Project: AVN DDBMS (Nigerian Army Aviation Dashboard Frontend)')
doc.add_paragraph(f'Workspace: {root}')
doc.add_paragraph(f'Report Date: {date.today().isoformat()}')
doc.add_paragraph('Platform Targets: Web (Vite/React), Desktop (Tauri + Electron configs present)')

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    'This project is a modular, enterprise-scale operational dashboard for Nigerian Army Aviation. '
    'It is built with React 19 + TypeScript and organized by features, supporting role-based access, '
    'real-time capabilities, finance workflows, inventory and procurement lifecycles, operations, '
    'maintenance tracking, document handling, reporting, messaging, and user/security administration.'
)

doc.add_heading('2. Technology Stack', level=1)
for item in [
    'Frontend framework: React 19 + TypeScript + Vite',
    'Routing: React Router with lazy-loaded route modules',
    'State/API: Redux Toolkit + RTK Query + Axios base query',
    'Styling: Tailwind CSS v4 + design tokens',
    'Utilities: html2canvas + jsPDF for PDF export, xlsx for spreadsheet workflows',
    'Realtime: Laravel Echo + Pusher',
    'Desktop packaging: Tauri 2 and Electron build paths',
    'Testing: Vitest + Testing Library',
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('3. Major Feature Areas', level=1)
features = {
    'Authentication': 'Login, registration, forgot-password, OTP verification flow, guarded routes.',
    'Finance': 'Expenses and budgets/accounts, budget statements, income recording, dynamic currency handling, expense edit/view slider and PDF export.',
    'Inventory & Procurement': 'Inventory CRUD views, requisitions, suppliers, RFQ details, purchase orders, procurement lifecycle.',
    'Operations': 'Operations list/detail/create/edit, after-action reporting, map/location workflows.',
    'Maintenance': 'Work orders, defects, planning/compliance/tracking flows, maintenance program CRUD.',
    'Assets': 'Asset listing/details and create/edit for aircraft, vehicles, facilities; defect workflows.',
    'Personnel & Units': 'Personnel directory/detail and unit management pages.',
    'Security & Users': 'Security module plus user details/add-user views, role/permission route gating.',
    'Reports & Documents': 'Dedicated modules for reporting and document management.',
    'Messaging': 'Message module with compose and realtime support scaffolding.',
}
for k, v in features.items():
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f'{k}: ').bold = True
    p.add_run(v)

doc.add_heading('4. Routing Coverage Snapshot', level=1)
for route in [
    '/auth/login, /auth/register, /auth/forgot-password, /auth/verify-otp',
    '/dashboard/finance, /dashboard/finance/add, /dashboard/finance/budgets/add, /dashboard/finance/budgets/:id, /dashboard/finance/budgets/:id/statement',
    '/dashboard/inventory and related requisition/purchase-order/rfq/procurement-lifecycle routes',
    '/dashboard/operations and add/edit/details/after-action-report routes',
    '/dashboard/maintenance and work-order/details/defects/compliance/tracking routes',
    '/dashboard/assets, /dashboard/personnel, /dashboard/units, /dashboard/documents, /dashboard/reports, /dashboard/messages, /dashboard/security, /dashboard/users',
]:
    doc.add_paragraph(route, style='List Bullet')

doc.add_heading('5. Build & Deployment Readiness', level=1)
for item in [
    'Production web build: tsc -b && vite build',
    'Desktop build: Tauri (macOS Intel x86_64 verified)',
    'Electron build scripts available for Windows packaging',
    'Firebase deploy workflow present in workspace operations',
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('6. Screenshots', level=1)
image_items = [
    ('Login Screen (/auth/login)', shots / 'login.png'),
    ('Registration Screen (/auth/register)', shots / 'register.png'),
    ('Forgot Password Screen (/auth/forgot-password)', shots / 'forgot-password.png'),
]
for title, path in image_items:
    doc.add_heading(title, level=2)
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.5))
        doc.add_paragraph(f'Captured from: {path.name}')
    else:
        doc.add_paragraph(f'Screenshot missing: {path}')

doc.add_heading('7. Notes', level=1)
doc.add_paragraph(
    'Authenticated internal module screenshots (finance dashboard, inventory, maintenance, etc.) require valid login credentials. '
    'This report includes all publicly reachable entry/auth screens captured from the running app.'
)

doc.save(out)
print(out)
